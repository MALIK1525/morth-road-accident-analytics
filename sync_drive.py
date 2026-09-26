"""Re-download the MoRTH workbook from Google Drive and rebuild data.json.
Usage: python sync_drive.py "<drive-file-id-or-url>"
Requires: pip install openpyxl
"""
import json, re, sys, urllib.request
from pathlib import Path
HERE = Path(__file__).parent
KEYS = {
 "State-Year Records": ["stateUt","year","accidents","fatalities","injured","zone"],
 "India-Year Records": ["year","accidents","fatalities","injured"],
 "Zone-Year Records": ["zone","year","accidents","fatalities","injured"],
 "State Regression Slopes": ["stateUt","nYears","slopePerYear","rSquared","meanAccidents","trend"],
 "ML Predictions": ["state","year","actual","baseline","ols","ridge","knn","rf","gb"],
 "ML Clustering": ["state","zone","clusterId","meanAccidents","meanFatalities","fatalityRate","slope","r2"],
 "Validation Checks": ["check","description","computedValue","publishedValue","difference","result"],
 "Source Provenance": ["sourceName","dataset","publisher","yearCoverage","verificationStatus","notes"],
}
NAMES = {"State-Year Records":"states","India-Year Records":"india","Zone-Year Records":"zones",
 "State Regression Slopes":"slopes","ML Predictions":"mlPred","ML Clustering":"clusters",
 "Validation Checks":"validation","Source Provenance":"sources","ML Model Comparison":"modelComp"}
def file_id(s):
    m = re.search(r"/d/([A-Za-z0-9_-]+)", s)
    return m.group(1) if m else s
def main(arg):
    fid = file_id(arg)
    url = "https://drive.google.com/uc?export=download&id=" + fid
    xlsx = HERE / "MoRTH_Complete_Data_Export.xlsx"
    urllib.request.urlretrieve(url, xlsx)
    import openpyxl
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    data = {}
    for sheet, cols in KEYS.items():
        ws = wb[sheet]; hdr = [c.value for c in ws[1]]
        rows = []
        for r in ws.iter_rows(min_row=2, values_only=True):
            if all(v is None for v in r): continue
            d = dict(zip(hdr, r))
            rows.append({k: d.get(k) for k in cols})
        data[NAMES[sheet]] = rows
    ws = wb["ML Model Comparison"]; hdr = [c.value for c in ws[1]]
    data["modelComp"] = [dict(zip(hdr, r)) for r in ws.iter_rows(min_row=2, values_only=True) if any(v is not None for v in r)]
    json.dump(data, open(HERE / "data.json", "w"))
    print("wrote data.json", {k: len(v) for k, v in data.items()})
if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "1SRdb_ImOcS7B8j5P856XG04Eftv3HSzJ")
